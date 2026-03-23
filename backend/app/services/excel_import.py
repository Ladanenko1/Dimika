from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Product, ProductVariant


def _normalize_key(value: str) -> str:
    return re.sub(r"[^0-9a-zа-я]+", "", value.lower().strip(), flags=re.IGNORECASE)


ALIASES_RAW: dict[str, list[str]] = {
    "article": ["Артикул", "article"],
    "name": [
        "Название товара",
        "Наименование товара",
        "Название",
        "Короткое наименование",
        "Полное название из каталога",
        "Расширенное наименование",
    ],
    "description": ["Описание", "Описание товара", "Расширенное наименование"],
    "product_type": ["Тип", "Тип изделия", "Под категория", "Раздел 3го уровня"],
    "category": ["Категория", "Товарная категория", "Раздел 2го уровня", "Раздел 1го уровня"],
    "color": ["Цвет", "цвет", "Цвет решетки"],
    "material": ["Материал", "Вид решетки", "Дизайн"],
    "country": ["Страна производителя", "Страна", "Страна происхождения"],
    "brand": ["Бренд", "Компания"],
    "length": ["Длина, см", "Длина", "Длина упаковки, мм"],
    "width": ["Ширина, см", "Ширина", "Ширина упаковки, мм"],
    "height": ["Высота, см", "Высота", "Высота упаковки, мм", "Глубина упаковки, мм"],
    "size": ["Размер", "Размер (длина, ширина, высота)", "Габариты"],
    "price": [
        "Цена (руб.)",
        "Итоговая цена со скидкой (руб.)",
        "РРЦ 11.12.25",
        "РРЦ по акции до 31.03.26",
        "РРЦ",
    ],
}

ALIASES: dict[str, list[str]] = {
    key: [_normalize_key(alias) for alias in aliases]
    for key, aliases in ALIASES_RAW.items()
}

MAX_ERRORS = 100


@dataclass
class ParsedCatalogRow:
    name: str
    product_type: str | None
    color: str | None
    category: str | None
    size: str | None
    material: str | None
    article: str
    brand: str | None
    country: str | None
    description: str | None
    price: Decimal


@dataclass
class ImportReport:
    total_rows: int = 0
    inserted: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


def _safe_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text


def _first_value(row: dict[str, Any], aliases: list[str]) -> str | None:
    for alias in aliases:
        value = _safe_str(row.get(alias))
        if value:
            return value
    return None


def _parse_decimal(value: str | None) -> Decimal | None:
    if value is None:
        return None
    sanitized = value.replace(" ", "").replace(",", ".")
    sanitized = re.sub(r"[^0-9.]+", "", sanitized)
    if not sanitized:
        return None
    try:
        parsed = Decimal(sanitized)
    except (InvalidOperation, ValueError):
        return None
    if parsed <= 0:
        return None
    return parsed.quantize(Decimal("0.01"))


def _extract_size(name: str | None, raw_size: str | None, length: str | None, width: str | None, height: str | None) -> str | None:
    if raw_size:
        return raw_size

    if length or width or height:
        parts = [length or "?", width or "?", height or "?"]
        return "x".join(parts)

    if name:
        match = re.search(
            r"(\d+(?:[.,]\d+)?)\s*[xх*]\s*(\d+(?:[.,]\d+)?)\s*[xх*]\s*(\d+(?:[.,]\d+)?)",
            name,
            flags=re.IGNORECASE,
        )
        if match:
            return "x".join([segment.replace(",", ".") for segment in match.groups()])
    return None


def _row_to_parsed(
    row: dict[str, Any],
    sheet_name: str,
    source_brand: str | None,
) -> ParsedCatalogRow | None:
    article = _first_value(row, ALIASES["article"])
    if article is None:
        return None

    name = _first_value(row, ALIASES["name"])
    description = _first_value(row, ALIASES["description"])
    if name is None:
        name = description
    if name is None:
        return None

    brand = _first_value(row, ALIASES["brand"]) or source_brand
    product_type = _first_value(row, ALIASES["product_type"])
    category = _first_value(row, ALIASES["category"]) or _safe_str(sheet_name)
    color = _first_value(row, ALIASES["color"])
    material = _first_value(row, ALIASES["material"])
    country = _first_value(row, ALIASES["country"])

    raw_size = _first_value(row, ALIASES["size"])
    length = _first_value(row, ALIASES["length"])
    width = _first_value(row, ALIASES["width"])
    height = _first_value(row, ALIASES["height"])
    size = _extract_size(name=name, raw_size=raw_size, length=length, width=width, height=height)

    raw_price = _first_value(row, ALIASES["price"])
    price = _parse_decimal(raw_price) or Decimal("1.00")

    return ParsedCatalogRow(
        name=name,
        product_type=product_type,
        color=color,
        category=category,
        size=size,
        material=material,
        article=article,
        brand=brand,
        country=country,
        description=description,
        price=price,
    )


def parse_supplier_excel(file_bytes: bytes, filename: str) -> tuple[list[ParsedCatalogRow], list[str]]:
    errors: list[str] = []
    rows: list[ParsedCatalogRow] = []
    source_brand = Path(filename).stem.strip() or None

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            try:
                header_row = next(iterator)
            except StopIteration:
                continue

            normalized_header: dict[int, str] = {}
            for idx, cell in enumerate(header_row):
                header_name = _safe_str(cell)
                if header_name:
                    normalized_header[idx] = _normalize_key(header_name)

            if not normalized_header:
                continue

            for row_idx, row_values in enumerate(iterator, start=2):
                if not row_values or not any(value is not None for value in row_values):
                    continue

                normalized_row: dict[str, Any] = {}
                for idx, value in enumerate(row_values):
                    header = normalized_header.get(idx)
                    if header is None:
                        continue
                    normalized_row[header] = value

                parsed = _row_to_parsed(normalized_row, sheet.title, source_brand)
                if parsed is None:
                    continue
                rows.append(parsed)
    finally:
        workbook.close()

    return rows, errors


def import_rows_to_db(
    db: Session,
    rows: list[ParsedCatalogRow],
    supplier_id: int | None = None,
) -> ImportReport:
    report = ImportReport(total_rows=len(rows))
    for idx, row in enumerate(rows, start=1):
        try:
            existing_variant = db.execute(
                select(ProductVariant).where(ProductVariant.article == row.article)
            ).scalar_one_or_none()

            if existing_variant:
                product = db.get(Product, existing_variant.product_id)
                if product is None:
                    raise RuntimeError(
                        f"Variant article '{row.article}' has no linked product"
                    )

                product.name = row.name
                product.type = row.product_type
                product.category = row.category
                product.brand = row.brand
                product.discription = row.description
                if supplier_id is not None:
                    product.supplier_id = supplier_id

                existing_variant.size = row.size
                existing_variant.color = row.color
                existing_variant.price = row.price
                existing_variant.material = row.material
                existing_variant.country = row.country
                report.updated += 1
            else:
                product = Product(
                    name=row.name,
                    discription=row.description,
                    brand=row.brand,
                    type=row.product_type,
                    category=row.category,
                    supplier_id=supplier_id,
                )
                variant = ProductVariant(
                    article=row.article,
                    size=row.size,
                    color=row.color,
                    price=row.price,
                    material=row.material,
                    country=row.country,
                )
                product.variants.append(variant)
                db.add(product)
                report.inserted += 1

            db.commit()
        except Exception as exc:  # noqa: BLE001
            db.rollback()
            report.skipped += 1
            if len(report.errors) < MAX_ERRORS:
                report.errors.append(f"row #{idx} article '{row.article}': {exc}")

    return report
